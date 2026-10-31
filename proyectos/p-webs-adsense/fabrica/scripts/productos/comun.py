"""Utilidades comunes de los productos digitales de pago (proyecto P, familia Claro).

Cada producto se genera con un script reproducible que lee los parámetros oficiales de params/<pais>/<año>/*.json
y los colores de sitios/<sitio>/sitio.json. Ninguna cifra legal se escribe a mano: las hojas «Parámetros» de los
Excel se rellenan desde esos archivos, con su fuente, y las fórmulas las leen por nombre.

Compatibilidad: solo funciones de Excel 2007 (SUMPRODUCT, SUMIFS, INDEX/MATCH, EDATE, EOMONTH, WORKDAY, PMT, RATE,
NPER, IRR…), que también funcionan igual en LibreOffice y Google Sheets. Nada de matrices dinámicas ni XLOOKUP.
"""
from __future__ import annotations

import json
import os
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule

FABRICA = Path(__file__).resolve().parents[2]
SALIDA = FABRICA / 'productos'
FUENTE = 'Arial'  # presente en Windows, macOS, LibreOffice y Google Sheets


def leer_json(rel: str):
    return json.loads((FABRICA / rel).read_text(encoding='utf-8'))


def sitio(id_: str) -> dict:
    return leer_json(f'sitios/{id_}/sitio.json')


def params(pais: str, anio: int, bloque: str) -> dict:
    return leer_json(f'params/{pais}/{anio}/{bloque}.json')


def producto(id_: str) -> dict:
    for p in leer_json('productos.json')['productos']:
        if p['id'] == id_:
            return p
    raise KeyError(id_)


def url_publica(sitio_id: str) -> str:
    """Dirección pública de hoy (familia.json) o, si no hay, el dominio previsto del sitio."""
    fam = leer_json('familia.json')
    for s in fam['sitios']:
        if s['id'] == sitio_id and s.get('url'):
            return s['url'].rstrip('/')
    return sitio(sitio_id)['dominio'].rstrip('/')


def hexc(c: str) -> str:
    return c.lstrip('#').upper()


def mezclar(c: str, blanco: float) -> str:
    """Mezcla un color con blanco (0 = el color, 1 = blanco)."""
    c = hexc(c)
    r, g, b = (int(c[i:i + 2], 16) for i in (0, 2, 4))
    f = lambda x: round(x + (255 - x) * blanco)
    return f'{f(r):02X}{f(g):02X}{f(b):02X}'


class Estilo:
    """Paleta y estilos de celda de un sitio."""

    def __init__(self, sitio_id: str, moneda: str = 'eur'):
        s = sitio(sitio_id)
        c = s['colores']
        self.sitio = s
        self.nombre = s['nombre']
        self.pri = hexc(c['primario'])
        self.osc = hexc(c['primarioOscuro'])
        self.marca = hexc(c['marca'])
        self.acento = hexc(c['acento'])
        self.tinta = hexc(c['tinta'])
        self.suave = hexc(c['suave'])
        self.borde = hexc(c['borde'])
        self.tenue = hexc(c['tenue'])
        self.entrada = mezclar(c['acento'], 0.72)  # amarillo claro: celdas que rellena el usuario
        self.moneda = moneda
        self.fmt_dinero = '#,##0.00 "€"' if moneda == 'eur' else '"$"#,##0.00'
        self.fmt_dinero0 = '#,##0 "€"' if moneda == 'eur' else '"$"#,##0'
        self.fmt_fecha = 'dd/mm/yyyy' if moneda == 'eur' else 'mm/dd/yyyy'
        self.fmt_pct = '0.00" %"' if moneda == 'eur' else '0.00"%"'
        fino = Side(style='thin', color=self.borde)
        self.b_fino = Border(left=fino, right=fino, top=fino, bottom=fino)
        m = Side(style='thin', color=self.marca)
        self.b_entrada = Border(left=m, right=m, top=m, bottom=m)
        p = Side(style='medium', color=self.pri)
        self.b_resultado = Border(left=p, right=p, top=p, bottom=p)

    # --- fuentes y rellenos
    def f(self, size=10, bold=False, color=None, italic=False, underline=None):
        return Font(name=FUENTE, size=size, bold=bold, color=color or self.tinta, italic=italic, underline=underline)

    def relleno(self, c):
        return PatternFill('solid', start_color=c, end_color=c)


def nuevo_libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    # Que Excel recalcule todo al abrir (el archivo se guarda sin valores en caché)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    return wb


def nombre(wb: Workbook, nom: str, hoja: str, celda: str):
    """Nombre definido (Excel, LibreOffice y Google Sheets lo importan como intervalo con nombre)."""
    col, fila = re.match(r'([A-Z]+)(\d+)', celda).groups()
    ref = f"'{hoja}'!${col}${fila}"
    wb.defined_names[nom] = DefinedName(nom, attr_text=ref)


def nombre_rango(wb: Workbook, nom: str, hoja: str, rango: str):
    a, b = rango.split(':')
    ca, fa = re.match(r'([A-Z]+)(\d+)', a).groups()
    cb, fb = re.match(r'([A-Z]+)(\d+)', b).groups()
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{hoja}'!${ca}${fa}:${cb}${fb}")


class Hoja:
    """Ayudante para escribir una hoja con el estilo del sitio."""

    def __init__(self, wb: Workbook, titulo: str, E: Estilo, color_pestana: str | None = None, anchos: dict | None = None):
        self.ws = wb.create_sheet(titulo)
        self.wb = wb
        self.E = E
        self.titulo = titulo
        self.ws.sheet_properties.tabColor = color_pestana or E.pri
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_view.zoomScale = 100
        for col, w in (anchos or {}).items():
            self.ws.column_dimensions[col].width = w
        ps = self.ws.page_setup
        ps.orientation = 'landscape'
        ps.paperSize = self.ws.PAPERSIZE_A4 if E.moneda == 'eur' else self.ws.PAPERSIZE_LETTER
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        self.ws.print_options.horizontalCentered = True
        self.ws.page_margins.left = self.ws.page_margins.right = 0.4
        self.ws.page_margins.top = self.ws.page_margins.bottom = 0.5
        self.ws.oddFooter.center.text = f'{E.nombre} · &A · &P/&N'
        self.ws.oddFooter.center.size = 8

    def c(self, ref, valor=None, *, fmt=None, font=None, fill=None, border=None, al=None, wrap=False, h=None, v='center'):
        cel = self.ws[ref]
        if valor is not None:
            cel.value = valor
        cel.font = font or self.E.f()
        if fmt:
            cel.number_format = fmt
        if fill:
            cel.fill = self.E.relleno(fill)
        if border:
            cel.border = border
        cel.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
        return cel

    def cabecera(self, titulo: str, sub: str = '', ancho_col='H'):
        """Banda de título con el nombre del sitio."""
        E = self.E
        ws = self.ws
        ws.row_dimensions[1].height = 8
        ws.row_dimensions[2].height = 34
        ws.row_dimensions[3].height = 20
        for fila in (1, 2, 3, 4):
            for col in range(1, _col(ancho_col) + 1):
                ws.cell(fila, col).fill = E.relleno(E.osc if fila < 4 else E.pri)
        ws.row_dimensions[4].height = 4
        self.c('B2', titulo, font=E.f(18, True, 'FFFFFF'))
        self.c('B3', sub, font=E.f(10, False, mezclar(E.marca, 0.55)))
        self.c(f'{ancho_col}2', E.nombre, font=E.f(11, True, E.acento), h='right')
        return 6

    def seccion(self, fila: int, texto: str, desde='B', hasta='H'):
        E = self.E
        for col in range(_col(desde), _col(hasta) + 1):
            self.ws.cell(fila, col).fill = E.relleno(E.tenue)
            self.ws.cell(fila, col).border = Border(bottom=Side(style='medium', color=E.marca))
        self.c(f'{desde}{fila}', texto, font=E.f(12, True, E.pri), fill=E.tenue, border=Border(bottom=Side(style='medium', color=E.marca)))
        self.ws.row_dimensions[fila].height = 22

    def etiqueta(self, ref, texto, ayuda: str | None = None, bold=False):
        self.c(ref, texto, font=self.E.f(10, bold), wrap=True)
        if ayuda:
            col, fila = re.match(r'([A-Z]+)(\d+)', ref).groups()
            # la ayuda va dos columnas a la derecha de la etiqueta (después del valor)

    def entrada(self, ref, valor=None, fmt=None, h='right'):
        cel = self.c(ref, valor, fmt=fmt, font=self.E.f(10, True, '1F3B57'), fill=self.E.entrada, border=self.E.b_entrada, h=h)
        cel.protection = Protection(locked=False)
        return cel

    def formula(self, ref, f, fmt=None, bold=False, fill=None, h='right', color=None):
        return self.c(ref, f, fmt=fmt, font=self.E.f(10, bold, color), fill=fill, border=self.E.b_fino, h=h)

    def resultado(self, ref, f, fmt=None, size=12):
        return self.c(ref, f, fmt=fmt, font=self.E.f(size, True, self.E.pri), fill=self.E.tenue, border=self.E.b_resultado, h='right')

    def nota(self, ref, texto, size=9, color=None, italic=False):
        return self.c(ref, texto, font=self.E.f(size, False, color or self.E.suave, italic=italic), wrap=True, v='top')

    def tabla_cab(self, fila: int, desde_col: int, textos: list[str], alto=30):
        for k, t in enumerate(textos):
            cel = self.ws.cell(fila, desde_col + k, t)
            cel.font = self.E.f(9, True, 'FFFFFF')
            cel.fill = self.E.relleno(self.E.pri)
            cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cel.border = self.E.b_fino
        self.ws.row_dimensions[fila].height = alto

    def leyenda(self, fila: int, col='B'):
        """Una sola celda amarilla con la regla del libro: lo amarillo lo rellenas tú."""
        E = self.E
        cel = self.c(f'{col}{fila}', '✎  Las celdas amarillas las rellenas tú; el resto se calcula solo.', font=E.f(9, True, '1F3B57'), fill=E.entrada, border=E.b_entrada)
        cel.alignment = Alignment(vertical='center', wrap_text=False)
        self.ws.row_dimensions[fila].height = 20

    def lista(self, rango: str, opciones: list[str] | str, error='Elige una opción de la lista.'):
        formula = opciones if isinstance(opciones, str) else '"' + ','.join(opciones) + '"'
        dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle='Valor no válido', error=error)
        self.ws.add_data_validation(dv)
        dv.add(rango)

    def numero(self, rango: str, minimo=0, maximo=None, decimal=True):
        dv = DataValidation(type='decimal' if decimal else 'whole', operator='between' if maximo is not None else 'greaterThanOrEqual',
                            formula1=str(minimo), formula2=str(maximo) if maximo is not None else None, allow_blank=True,
                            showErrorMessage=True, errorTitle='Número no válido', error='Escribe un número (sin el símbolo de moneda).')
        self.ws.add_data_validation(dv)
        dv.add(rango)

    def fecha(self, rango: str):
        dv = DataValidation(type='date', operator='greaterThan', formula1='1', allow_blank=True, showErrorMessage=True,
                            errorTitle='Fecha no válida', error='Escribe una fecha, por ejemplo 15/03/2027.' if self.E.moneda == 'eur' else 'Escribe una fecha, por ejemplo 04/15/2026.')
        self.ws.add_data_validation(dv)
        dv.add(rango)

    def proteger(self):
        """Protege las fórmulas sin contraseña (se quita en Revisar → Desproteger hoja)."""
        p = self.ws.protection
        p.sheet = True
        p.formatColumns = False
        p.formatRows = False
        p.selectLockedCells = False
        p.selectUnlockedCells = False

    def resaltar_si(self, rango: str, formula: str, color: str, bold=False, font_color=None):
        self.ws.conditional_formatting.add(rango, FormulaRule(formula=[formula], fill=self.E.relleno(color), font=Font(name=FUENTE, bold=bold, color=font_color) if (bold or font_color) else None))


def _col(letra: str) -> int:
    n = 0
    for ch in letra:
        n = n * 26 + (ord(ch) - 64)
    return n


def _letra(n: int) -> str:
    return get_column_letter(n)


def portada(wb: Workbook, E: Estilo, *, titulo: str, subtitulo: str, edicion: str, contenido: list[tuple[str, str]],
            avisos: list[str], url: str, pasos: list[str]):
    """Primera hoja: portada con índice enlazado, tres pasos y avisos legales."""
    h = Hoja(wb, 'Portada', E, E.acento, {'A': 3, 'B': 4, 'C': 34, 'D': 60, 'E': 4})
    ws = h.ws
    for fila in range(1, 13):
        ws.row_dimensions[fila].height = 18
        for col in range(1, 6):
            ws.cell(fila, col).fill = E.relleno(E.osc)
    h.c('C3', E.nombre.upper(), font=E.f(11, True, E.acento))
    h.c('C5', titulo, font=E.f(26, True, 'FFFFFF'))
    ws.row_dimensions[5].height = 40
    h.c('C7', subtitulo, font=E.f(12, False, mezclar(E.marca, 0.6)))
    ws.merge_cells('C7:D8')
    ws['C7'].alignment = Alignment(wrap_text=True, vertical='top')
    h.c('C10', edicion, font=E.f(10, True, E.acento))
    for col in range(1, 6):
        ws.cell(12, col).fill = E.relleno(E.pri)
    ws.row_dimensions[12].height = 6
    fila = 14
    h.c(f'C{fila}', 'Empieza en tres pasos', font=E.f(13, True, E.pri))
    fila += 1
    for k, p in enumerate(pasos, 1):
        h.c(f'B{fila}', k, font=E.f(12, True, 'FFFFFF'), fill=E.marca, h='center')
        h.c(f'C{fila}', p, font=E.f(10), wrap=True)
        ws.merge_cells(f'C{fila}:D{fila}')
        ws.row_dimensions[fila].height = 30
        fila += 1
    fila += 1
    h.c(f'C{fila}', 'Qué hay en este libro', font=E.f(13, True, E.pri))
    fila += 1
    for hoja, desc in contenido:
        cel = h.c(f'C{fila}', f'→ {hoja}', font=E.f(10, True, E.pri, underline='single'))
        cel.hyperlink = f"#'{hoja}'!A1"
        h.c(f'D{fila}', desc, font=E.f(10, False, E.suave), wrap=True)
        ws.row_dimensions[fila].height = 28 if len(desc) > 70 else 18
        fila += 1
    fila += 1
    h.c(f'C{fila}', 'Antes de usarlo', font=E.f(13, True, E.pri))
    fila += 1
    for a in avisos:
        h.c(f'C{fila}', a, font=E.f(9, False, E.suave), wrap=True, v='top')
        ws.merge_cells(f'C{fila}:D{fila}')
        ws.row_dimensions[fila].height = 15 * max(1, (len(a) // 95) + 1)
        fila += 1
    fila += 1
    cel = h.c(f'C{fila}', url, font=E.f(10, True, E.pri, underline='single'))
    cel.hyperlink = url
    h.c(f'D{fila}', 'Contacto: hola@numimates.com', font=E.f(10, False, E.suave), h='right')
    ws.sheet_view.zoomScale = 110
    return h


def instrucciones(wb: Workbook, E: Estilo, titulo: str, bloques: list[tuple[str, list[str]]]):
    h = Hoja(wb, 'Instrucciones', E, E.marca, {'A': 3, 'B': 4, 'C': 110})
    fila = h.cabecera(titulo, 'Léelo una vez: son cinco minutos', 'C')
    for t, lineas in bloques:
        h.seccion(fila, t, 'B', 'C')
        fila += 1
        for ln in lineas:
            h.c(f'B{fila}', '•', font=h.E.f(10, True, h.E.marca), h='center', v='top')
            h.c(f'C{fila}', ln, font=h.E.f(10), wrap=True, v='top')
            h.ws.row_dimensions[fila].height = 15 * max(1, (len(ln) // 120) + 1)
            fila += 1
        fila += 1
    return h


def hoja_parametros(wb: Workbook, E: Estilo, filas: list[dict], titulo='Parámetros', sub=''):
    """Hoja de parámetros: cada fila es {clave, concepto, valor, fmt, fuente, url}. Crea un nombre definido por clave."""
    h = Hoja(wb, titulo, E, E.suave, {'A': 3, 'B': 46, 'C': 16, 'D': 58, 'E': 3})
    fila = h.cabecera(titulo, sub or 'Cifras oficiales que usan las fórmulas. Salen de los archivos de parámetros del sitio, con su fuente.', 'D')
    h.tabla_cab(fila, 2, ['Concepto', 'Valor', 'Fuente oficial'])
    fila += 1
    for p in filas:
        if p.get('seccion'):
            h.seccion(fila, p['seccion'], 'B', 'D')
            fila += 1
            continue
        h.c(f'B{fila}', p['concepto'], wrap=True, border=E.b_fino)
        h.formula(f'C{fila}', p['valor'], fmt=p.get('fmt'))
        cel = h.c(f'D{fila}', p.get('fuente', ''), font=E.f(8, False, E.pri if p.get('url') else E.suave, underline='single' if p.get('url') else None), wrap=True, border=E.b_fino)
        if p.get('url'):
            cel.hyperlink = p['url']
        h.ws.row_dimensions[fila].height = 24 if len(p.get('fuente', '')) > 60 or len(p['concepto']) > 50 else 16
        if p.get('clave'):
            nombre(wb, p['clave'], titulo, f'C{fila}')
        fila += 1
    h.proteger()
    return h


def fuente_de(bloque: dict) -> tuple[str, str]:
    f = bloque.get('fuente') or {}
    if isinstance(f, dict):
        return f.get('nombre', ''), f.get('url', '')
    return str(f), ''


def guardar(wb: Workbook, ruta: Path):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    # La portada se abre primero
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == wb.worksheets[0].title
    wb.properties.creator = 'Familia Claro'
    wb.properties.lastModifiedBy = 'Familia Claro'
    wb.save(ruta)


def iso(d: date) -> str:
    return d.isoformat()


DIAS_ES = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
MESES_ES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def fecha_larga(d: date) -> str:
    return f'{DIAS_ES[d.weekday()]} {d.day} de {MESES_ES[d.month - 1]} de {d.year}'


def fecha_iso(s: str) -> date:
    return date.fromisoformat(s)


# ---------------------------------------------------------------- PDF (HTML que Playwright imprime)

def _b64(p: Path) -> str:
    import base64
    return base64.b64encode(p.read_bytes()).decode()


def css_pdf(E: Estilo) -> str:
    F = FABRICA / 'node_modules/@fontsource'
    fuentes = ''
    for fam, peso, arch in [('Manrope', 700, 'manrope/files/manrope-latin-700-normal.woff2'), ('Manrope', 800, 'manrope/files/manrope-latin-800-normal.woff2'),
                            ('Inter', 400, 'inter/files/inter-latin-400-normal.woff2'), ('Inter', 600, 'inter/files/inter-latin-600-normal.woff2'),
                            ('Inter', 700, 'inter/files/inter-latin-700-normal.woff2')]:
        fuentes += f"@font-face{{font-family:{fam};font-weight:{peso};src:url(data:font/woff2;base64,{_b64(F / arch)}) format('woff2')}}"
    # latin-ext para ç, l·l, ñ ya está en latin; € también
    return fuentes + f"""
@page{{size:{'A4' if E.moneda == 'eur' else 'Letter'};margin:16mm 16mm 18mm}}
*{{box-sizing:border-box}}
html{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
body{{font-family:Inter,Arial,sans-serif;font-size:10.5pt;line-height:1.5;color:#{E.tinta};margin:0}}
h1,h2,h3{{font-family:Manrope,Arial,sans-serif;line-height:1.2;color:#{E.osc}}}
h1{{font-size:24pt;margin:0 0 6pt;font-weight:800}}
h2{{font-size:15pt;margin:20pt 0 6pt;padding-bottom:4pt;border-bottom:2pt solid #{E.marca};break-after:avoid}}
h3{{font-size:11.5pt;margin:14pt 0 4pt;break-after:avoid}}
p{{margin:0 0 7pt}} ul,ol{{margin:0 0 8pt;padding-left:16pt}} li{{margin:0 0 3pt}}
a{{color:#{E.pri}}}
.portada{{height:247mm;display:flex;flex-direction:column;justify-content:space-between;background:#{E.osc};color:#fff;margin:-16mm -16mm 0;padding:24mm 20mm;break-after:page}}
.portada .marca{{font-family:Manrope;font-weight:800;color:#{E.acento};letter-spacing:.08em;font-size:11pt}}
.portada h1{{color:#fff;font-size:34pt;margin:14mm 0 6mm}}
.portada .sub{{font-size:14pt;color:#{mezclar(E.marca, 0.55)};max-width:150mm}}
.portada .chips{{display:flex;flex-wrap:wrap;gap:6pt;margin-top:10mm}}
.portada .chips span{{border:1pt solid #{E.marca};border-radius:20pt;padding:4pt 10pt;font-size:9.5pt;color:#fff}}
.portada .pie{{font-size:9pt;color:#{mezclar(E.marca, 0.55)};border-top:1pt solid #{E.pri};padding-top:6mm}}
.caja{{border:1pt solid #{E.borde};border-left:4pt solid #{E.marca};background:#{E.tenue};border-radius:4pt;padding:9pt 12pt;margin:10pt 0;break-inside:avoid}}
.caja.aviso{{border-left-color:#{E.acento};background:#FFFBEF}}
.carta{{border:1pt solid #{E.borde};border-radius:4pt;padding:14pt 16pt;margin:8pt 0 12pt;font-size:10pt;background:#fff;break-inside:avoid}}
.carta p{{margin:0 0 6pt}}
.hueco{{background:#{E.entrada};padding:0 3pt;border-radius:2pt}}
table{{border-collapse:collapse;width:100%;margin:6pt 0 10pt;font-size:9.5pt;break-inside:avoid}}
th{{background:#{E.pri};color:#fff;text-align:left;padding:5pt 7pt;font-weight:600}}
td{{border-bottom:1pt solid #{E.borde};padding:5pt 7pt;vertical-align:top}}
.check li{{list-style:none;margin-left:-14pt}} .check li::before{{content:'☐';margin-right:6pt;color:#{E.pri};font-size:11pt}}
.pasos{{counter-reset:p;list-style:none;padding:0}} .pasos>li{{counter-increment:p;position:relative;padding-left:26pt;margin-bottom:8pt}}
.pasos>li::before{{content:counter(p);position:absolute;left:0;top:0;width:18pt;height:18pt;border-radius:50%;background:#{E.marca};color:#fff;font-weight:700;text-align:center;line-height:18pt;font-size:9pt}}
.fuentes{{font-size:8.5pt;color:#{E.suave}}}
.legal{{font-size:8.5pt;color:#{E.suave};border-top:1pt solid #{E.borde};margin-top:16pt;padding-top:8pt}}
.dos{{display:grid;grid-template-columns:1fr 1fr;gap:10pt}}
"""


def html_pdf(E: Estilo, titulo: str, cuerpo: str) -> str:
    return f'<!doctype html><html lang="{"es" if E.moneda == "eur" else "es-US"}"><head><meta charset="utf-8"><title>{titulo}</title><style>{css_pdf(E)}</style></head><body>{cuerpo}</body></html>'


def portada_pdf(E: Estilo, titulo: str, sub: str, chips: list[str], pie: str) -> str:
    ch = ''.join(f'<span>{c}</span>' for c in chips)
    return f'<section class="portada"><div><div class="marca">{E.nombre.upper()}</div><h1>{titulo}</h1><p class="sub">{sub}</p><div class="chips">{ch}</div></div><div class="pie">{pie}</div></section>'


def escribir(ruta: Path, texto: str):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(texto, encoding='utf-8')


def celda_nombre(wb, nom: str):
    """(hoja, celda) de un nombre definido de una sola celda."""
    ref = wb.defined_names[nom].attr_text
    hoja, celda = ref.rsplit('!', 1)
    return wb[hoja.strip("'")], celda.replace('$', '')


def poner(wb, nom: str, valor):
    ws, c = celda_nombre(wb, nom)
    ws[c] = valor
