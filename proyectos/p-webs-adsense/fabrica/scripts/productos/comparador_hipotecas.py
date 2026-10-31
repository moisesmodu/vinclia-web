"""Comparador de hipotecas para negociar (HogarClaro): hasta 4 ofertas (FEIN) con TAE aproximada, coste total,
vinculaciones, escenarios de Euríbor, amortización anticipada y cuadro de amortización. Más un PDF de preguntas.

    python3 scripts/productos/comparador_hipotecas.py [--demo]

Cifras de params/es/<año>/vivienda.json (Euríbor, topes de comisión de la Ley 5/2019). Las fórmulas reproducen
cuadro(), hipotecaFija(), hipotecaVariable(), tae() y amortizacionAnticipada() de src/lib/calc/vivienda.ts.
"""
from __future__ import annotations

import sys

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment

from comun import (Estilo, Hoja, SALIDA, escribir, fuente_de, guardar, hoja_parametros, html_pdf, instrucciones, mezclar, nombre,
                   nombre_rango, nuevo_libro, params, portada, portada_pdf, producto, url_publica, poner, leer_json, MESES_ES)

ID = 'comparador-hipotecas'
ANIO = leer_json('sitios/vivienda/sitio.json')['anioPorDefecto']
V = params('es', ANIO, 'vivienda')
MAXM = 480  # 40 años
OF = 'DEFG'  # columnas de las 4 ofertas
AVISO = 'Cálculo orientativo para comparar ofertas: la TAE oficial de la FEIN (Orden EHA/2899/2011) puede diferir unas centésimas. No es asesoramiento financiero.'


def ultimo_euribor():
    serie = [x for x in V['euribor_12m']['serie'] if x.get('valor') is not None]
    return serie[-1]


def construir(demo=False):
    E = Estilo('vivienda')
    P = producto(ID)
    wb = nuevo_libro()
    url = url_publica('vivienda')
    eu = ultimo_euribor()
    mes_eu = f'{MESES_ES[int(eu["mes"][5:]) - 1]} de {eu["mes"][:4]}'
    C = V['comision_amortizacion_anticipada']
    fc = fuente_de(C)
    fe = fuente_de(V['euribor_12m'])
    crit = V['criterios_editoriales_no_normativos']

    portada(wb, E, titulo=P['nombre']['es'], subtitulo='Pon hasta cuatro ofertas (FEIN) una al lado de otra y mira lo que de verdad te cuesta cada una: TAE con seguros, vinculaciones, escenarios de Euríbor y amortizar antes.',
            edicion=f'Edición {ANIO} · Euríbor de {mes_eu}: {str(eu["valor"]).replace(".", ",")} %',
            pasos=['En «Tu hipoteca» pon el importe, el plazo y el escenario de Euríbor.',
                   'En «Ofertas» copia los datos de cada FEIN: tipo, TIN, diferencial, comisiones, seguros y productos vinculados.',
                   'Mira «Comparativa»: coste total, TAE real y cuál conviene. Lleva al banco las preguntas del PDF.'],
            contenido=[('Comparativa', 'Las cuatro ofertas resumidas, ordenadas por coste total, con gráfico'), ('Tu hipoteca', 'Importe, plazo y escenario de Euríbor'),
                       ('Ofertas', 'Los datos de cada FEIN y sus resultados'), ('Escenarios Euríbor', 'Cuánto pagarías si el Euríbor baja o sube'),
                       ('Amortizar antes', 'Reducir cuota o plazo: ahorro y comisión legal'), ('Cuadro anual', 'Cuadro de amortización por años de la oferta que elijas'),
                       ('Cuadro mensual', 'Mes a mes de las cuatro ofertas (lo usan las demás hojas)'), ('Parámetros', 'Euríbor y topes legales con su fuente')],
            avisos=[AVISO, f'El Euríbor por defecto es la media de {mes_eu} ({eu["valor"]} %). Cámbialo o elige un escenario: nadie sabe cuánto valdrá.',
                    'Contenido preparado con ayuda de inteligencia artificial y revisado con las fuentes oficiales citadas (Ley 5/2019, Banco de España, BOE).',
                    'Funciona en Excel, LibreOffice y Google Sheets. Las hojas de cálculo están protegidas sin contraseña: Revisar → Desproteger hoja.'],
            url=url + '/' + P['ruta']['es'] + '/')
    instrucciones(wb, E, 'Cómo se usa el comparador', [
        ('1. Pide la FEIN', ['La Ficha Europea de Información Normalizada es obligatoria y gratuita: el banco te la tiene que dar con las condiciones personalizadas y al menos 10 días naturales antes de firmar (Ley 5/2019, art. 14). Pídela a cada banco aunque solo estés comparando.',
                             'De la FEIN necesitas: tipo de interés (TIN) con y sin bonificaciones, diferencial si es variable, años a tipo fijo si es mixta, comisiones, gastos de tasación y el precio de los seguros y productos que te exigen para bonificar.']),
        ('2. Rellena las ofertas', ['Fija: el TIN de todo el plazo. Mixta: el TIN de los años fijos, cuántos años son y el diferencial sobre el Euríbor del resto. Variable: el tipo del primer año (pon 1 en años iniciales) y el diferencial.',
                                    'Seguros: el precio anual de vida y hogar si los contratas con el banco para bonificar. Otros vinculados: lo que cuestan al mes la cuenta, la tarjeta, el plan de pensiones o la nómina que te piden.',
                                    'TIN o diferencial sin bonificar: el que tendrías sin contratar nada. El comparador te dice si bonificarte compensa.']),
        ('3. Lee los resultados', ['Coste total = intereses + seguros y vinculaciones + gastos iniciales. Es la cifra para comparar. La TAE aproximada incluye las vinculaciones que pongas, así que puede ser más alta que la de la FEIN.',
                                   'En las variables y mixtas, la cuota después del periodo inicial depende del Euríbor: mira «Escenarios Euríbor» para ver cuánto subiría.']),
        ('4. Amortizar antes', ['En «Amortizar antes» elige la oferta, el año y el importe: verás cuánto ahorras reduciendo cuota o plazo y la comisión máxima que permite la ley.']),
        ('Qué no hace', ['No incluye la subida de las primas con la edad, la revisión de los seguros ni los cambios de condiciones a mitad del préstamo. ' + AVISO]),
        ('Ayuda', ['¿Algo no cuadra? Escribe a hola@numimates.com.']),
    ])
    hoja_parametros(wb, E, [
        {'seccion': 'Euríbor a 12 meses'},
        {'clave': 'EURIBOR_ULTIMO', 'concepto': f'Media mensual de {mes_eu} (%)', 'valor': eu['valor'], 'fmt': '0.000', 'fuente': fe[0], 'url': fe[1]},
        {'seccion': 'Comisión por amortización anticipada (topes legales)'},
        {'clave': 'COM_FIJO_10', 'concepto': 'Tipo fijo: primeros 10 años (% del capital amortizado)', 'valor': C['fijo']['primeros_10_anios_pct'], 'fmt': '0.00', 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_FIJO_RESTO', 'concepto': 'Tipo fijo: después de 10 años (%)', 'valor': C['fijo']['resto_pct'], 'fmt': '0.00', 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_VAR_A', 'concepto': 'Variable, opción A: tope (%)', 'valor': C['variable']['opcion_a']['tope_pct'], 'fmt': '0.00', 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_VAR_A_ANIOS', 'concepto': 'Variable, opción A: durante los primeros (años)', 'valor': C['variable']['opcion_a']['durante_anios'], 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_VAR_B', 'concepto': 'Variable, opción B: tope (%)', 'valor': C['variable']['opcion_b']['tope_pct'], 'fmt': '0.00', 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_VAR_B_ANIOS', 'concepto': 'Variable, opción B: durante los primeros (años)', 'valor': C['variable']['opcion_b']['durante_anios'], 'fuente': fc[0], 'url': fc[1]},
        {'clave': 'COM_VAR_DESPUES', 'concepto': 'Variable: después (%)', 'valor': C['variable']['despues'], 'fmt': '0.00', 'fuente': fc[0], 'url': fc[1]},
        {'seccion': 'Criterio de los bancos (no es una norma)'},
        {'clave': 'ESFUERZO', 'concepto': 'Esfuerzo máximo habitual: cuota sobre ingresos netos (%)', 'valor': crit['esfuerzo_maximo_cuota_sobre_ingresos_netos_pct'], 'fuente': 'Práctica habitual de la banca; no lo fija ninguna norma'},
    ], sub=f'Cifras de params/es/{ANIO}/vivienda.json a {V["actualizado"]}.')

    # ------------------------------------------------------------ Tu hipoteca
    ht = Hoja(wb, 'Tu hipoteca', E, E.acento, {'A': 3, 'B': 50, 'C': 22, 'D': 66})
    fila = ht.cabecera('Tu hipoteca', 'Lo común a todas las ofertas', 'D')
    ht.leyenda(fila); fila += 2
    ESC = ['Euríbor de hoy', 'Baja 1 punto', 'Sube 1 punto', 'Sube 2 puntos', 'Personalizado']
    for texto, defecto, tipo, ayuda, nom in [
        ('Importe del préstamo (€)', 180000, 'dinero', 'Lo que te presta el banco (no el precio de la casa).', 'CAPITAL'),
        ('Plazo (años)', 30, 'anios', 'Hasta 40 años.', 'PLAZO'),
        ('Euríbor de partida (%)', '=EURIBOR_ULTIMO', 'pct', f'Último dato publicado: {mes_eu}. Cámbialo si quieres.', 'EURIBOR_HOY'),
        ('Escenario de Euríbor para después del periodo inicial', ESC[0], ESC, 'Se supone constante durante todo el préstamo.', 'ESCENARIO'),
        ('Euríbor personalizado (%)', 3.5, 'pct', 'Solo si eliges «Personalizado».', 'EURIBOR_PERS'),
        ('Tus ingresos netos al mes (opcional) (€)', 3200, 'dinero', '="Para ver el esfuerzo: los bancos suelen pedir que la cuota no pase del "&ESFUERZO&" %."', 'INGRESOS')]:
        ht.etiqueta(f'B{fila}', texto, bold=True)
        ht.entrada(f'C{fila}', defecto, fmt=E.fmt_dinero if tipo == 'dinero' else ('0.000' if tipo == 'pct' else None), h='left' if isinstance(tipo, list) else 'right')
        if isinstance(tipo, list):
            ht.lista(f'C{fila}', tipo)
        elif tipo == 'anios':
            ht.numero(f'C{fila}', 1, 40)
        ht.nota(f'D{fila}', ayuda)
        nombre(wb, nom, 'Tu hipoteca', f'C{fila}')
        ht.ws.row_dimensions[fila].height = 24
        fila += 1
    fila += 1
    ht.etiqueta(f'B{fila}', 'Euríbor que se usa en los cálculos (%)', bold=True)
    ht.resultado(f'C{fila}', '=CHOOSE(MATCH(ESCENARIO,{' + ','.join(f'"{x}"' for x in ESC) + '},0),EURIBOR_HOY,EURIBOR_HOY-1,EURIBOR_HOY+1,EURIBOR_HOY+2,EURIBOR_PERS)', fmt='0.000')
    nombre(wb, 'EURIBOR', 'Tu hipoteca', f'C{fila}')
    fila += 1
    ht.etiqueta(f'B{fila}', 'Meses del préstamo')
    ht.formula(f'C{fila}', '=MIN(480,ROUND(PLAZO*12,0))', fmt='0')
    nombre(wb, 'MESES', 'Tu hipoteca', f'C{fila}')
    fila += 2
    ht.nota(f'B{fila}', AVISO, italic=True)
    ht.ws.merge_cells(f'B{fila}:D{fila}')
    ht.proteger()

    # ------------------------------------------------------------ Ofertas (entradas y resultados)
    ho = Hoja(wb, 'Ofertas', E, E.pri, {'A': 3, 'B': 50, 'C': 3, 'D': 18, 'E': 18, 'F': 18, 'G': 18, 'H': 3})
    fila = ho.cabecera('Las ofertas', 'Una columna por FEIN', 'G')
    ho.leyenda(fila); fila += 2
    ho.tabla_cab(fila, 2, ['', '', 'Oferta 1', 'Oferta 2', 'Oferta 3', 'Oferta 4'])
    fila += 1
    O = {}
    demo_vals = {
        'banco': ['Banco Norte', 'Caja Sur', 'Banco Online', 'Mi banco'], 'tipo': ['Fija', 'Mixta', 'Variable', 'Fija'], 'tin': [2.6, 2.2, 2.5, 2.95],
        'anios_ini': [0, 5, 1, 0], 'dif': [0, 0.75, 0.6, 0], 'sin': [3.4, 1.1, 0.99, 2.95], 'apertura': [0, 0, 0, 0.5], 'tasacion': [350, 300, 0, 350],
        'otros': [0, 150, 0, 0], 'vida': [300, 280, 0, 0], 'hogar': [250, 220, 180, 0], 'vinc': [0, 6, 0, 0],
    }
    entradas = [
        ('banco', 'Banco u oferta', None, 'left'), ('tipo', 'Tipo de interés', ['Fija', 'Mixta', 'Variable'], 'center'),
        ('tin', 'TIN con bonificaciones (fija) o del periodo inicial (%)', '0.00', 'right'),
        ('anios_ini', 'Años del periodo inicial (mixta o variable)', '0', 'right'),
        ('dif', 'Diferencial sobre el Euríbor con bonificaciones (%)', '0.00', 'right'),
        ('sin', 'TIN (fija) o diferencial (variable) SIN bonificar (%)', '0.00', 'right'),
        ('apertura', 'Comisión de apertura (%)', '0.00', 'right'), ('tasacion', 'Tasación a tu cargo (€)', E.fmt_dinero0, 'right'),
        ('otros', 'Otros gastos a tu cargo (€)', E.fmt_dinero0, 'right'), ('vida', 'Seguro de vida vinculado (€ al año)', E.fmt_dinero0, 'right'),
        ('hogar', 'Seguro de hogar vinculado (€ al año)', E.fmt_dinero0, 'right'), ('vinc', 'Otros productos vinculados (€ al mes)', E.fmt_dinero, 'right'),
    ]
    for clave, texto, fmt, h in entradas:
        ho.etiqueta(f'B{fila}', texto, bold=True)
        for j, col in enumerate(OF):
            ho.entrada(f'{col}{fila}', demo_vals[clave][j] if (demo or clave == 'tipo') else None, fmt=None if isinstance(fmt, list) else fmt, h=h)
            nombre(wb, f'OF_{clave.upper()}_{j + 1}', 'Ofertas', f'{col}{fila}')
        if isinstance(fmt, list):
            ho.lista(f'D{fila}:G{fila}', fmt)
        O[clave] = fila
        fila += 1
    ho.etiqueta(f'B{fila}', 'Comisión por amortizar antes (%)', bold=True)
    for j, col in enumerate(OF):
        ho.entrada(f'{col}{fila}', f'=IF({col}{O["tipo"]}="Fija",COM_FIJO_10,COM_VAR_A)', fmt='0.00')
        nombre(wb, f'OF_COMISION_{j + 1}', 'Ofertas', f'{col}{fila}')
    ho.nota(f'H{fila}', '')
    O['comision'] = fila
    fila += 2
    ho.tabla_cab(fila, 2, ['Resultados', '', 'Oferta 1', 'Oferta 2', 'Oferta 3', 'Oferta 4'])
    fila += 1
    # Columnas del cuadro mensual de cada oferta (hoja «Cuadro mensual»): TIN, cuota, interés, capital, pendiente, flujo
    base_col = 4
    def col_cm(j, k):
        from openpyxl.utils import get_column_letter
        return get_column_letter(base_col + j * 6 + k)
    R0 = 8  # fila del mes 0 en «Cuadro mensual»
    rng = lambda j, k: f"'Cuadro mensual'!${col_cm(j, k)}${R0 + 1}:${col_cm(j, k)}${R0 + MAXM}"
    salidas = [
        ('cuota1', 'Cuota inicial (€ al mes)', lambda j, c: f"='Cuadro mensual'!{col_cm(j, 1)}{R0 + 1}", E.fmt_dinero, True),
        ('cuota2', 'Cuota después del periodo inicial (escenario)', lambda j, c: f"=IF({c}{O['tipo']}=\"Fija\",{c}{fila},INDEX({rng(j, 1)},MIN(MESES,{c}{O['anios_ini']}*12+1)))", E.fmt_dinero, False),
        ('vincmes', 'Seguros y vinculaciones (€ al mes)', lambda j, c: f'=({c}{O["vida"]}+{c}{O["hogar"]})/12+{c}{O["vinc"]}', E.fmt_dinero, False),
        ('gastos', 'Gastos iniciales a tu cargo', lambda j, c: f'=CAPITAL*{c}{O["apertura"]}/100+{c}{O["tasacion"]}+{c}{O["otros"]}', E.fmt_dinero0, False),
        ('intereses', 'Intereses de todo el préstamo', lambda j, c: f'=SUM({rng(j, 2)})', E.fmt_dinero0, False),
        ('vinctot', 'Seguros y vinculaciones de todo el plazo', lambda j, c: f'=VINC_MES_{j + 1}*MESES', E.fmt_dinero0, False),
        ('coste', 'COSTE TOTAL (intereses + vinculaciones + gastos)', lambda j, c: f'=INTERESES_{j + 1}+VINC_TOT_{j + 1}+GASTOS_{j + 1}', E.fmt_dinero0, True),
        ('tae', 'TAE aproximada con las vinculaciones', lambda j, c: f"=((1+IRR('Cuadro mensual'!${col_cm(j, 5)}${R0}:${col_cm(j, 5)}${R0 + MAXM},0.003))^12-1)*100", '0.00" %"', True),
        ('sinbon', 'Coste total sin bonificar (aprox., tipo constante)', lambda j, c: f'=PMT(IF({c}{O["tipo"]}="Fija",{c}{O["sin"]},MAX(0,EURIBOR+{c}{O["sin"]}))/1200,MESES,-CAPITAL)*MESES-CAPITAL+GASTOS_{j + 1}', E.fmt_dinero0, False),
        ('compensa', '¿Compensa bonificarte?', lambda j, c: f'=IF({c}{O["sin"]}="","—",IF(ABS(COSTE_{j + 1}-SINBON_{j + 1})<1,"Sin diferencia",IF(COSTE_{j + 1}<SINBON_{j + 1},"Sí, ahorras "&FIXED(SINBON_{j + 1}-COSTE_{j + 1},0)&" €","No: te cuesta "&FIXED(COSTE_{j + 1}-SINBON_{j + 1},0)&" € más")))', None, False),
        ('esfuerzo', 'Esfuerzo: cuota + vinculaciones sobre tus ingresos', lambda j, c: f'=IF(INGRESOS>0,(CUOTA1_{j + 1}+VINC_MES_{j + 1})/INGRESOS,"")', '0.0%', False),
        ('rank', 'Posición por coste total', lambda j, c: f'=IF(BANCO_OK_{j + 1},RANK(COSTE_{j + 1},RANK_RANGO,1),"")', '0"º"', True),
    ]
    for clave, texto, fs, fmt, fuerte in salidas:
        ho.etiqueta(f'B{fila}', texto, bold=fuerte)
        for j, col in enumerate(OF):
            f_ = fs(j, col)
            if clave == 'cuota2':
                f_ = f"=IF({col}{O['tipo']}=\"Fija\",CUOTA1_{j + 1},INDEX({rng(j, 1)},MIN(MESES,{col}{O['anios_ini']}*12+1)))"
            f_ = f'=IF(BANCO_OK_{j + 1},{f_[1:]},"")' if clave not in ('rank',) else f_
            (ho.resultado if fuerte else ho.formula)(f'{col}{fila}', f_, fmt=fmt, **({} if fuerte else {}))
            nombre(wb, f'{clave.upper()}_{j + 1}', 'Ofertas', f'{col}{fila}')
            if clave == 'vincmes':
                nombre(wb, f'VINC_MES_{j + 1}', 'Ofertas', f'{col}{fila}')
            if clave == 'vinctot':
                nombre(wb, f'VINC_TOT_{j + 1}', 'Ofertas', f'{col}{fila}')
        O[clave] = fila
        fila += 1
    nombre_rango(wb, 'RANK_RANGO', 'Ofertas', f'D{O["coste"]}:G{O["coste"]}')
    # Una oferta cuenta si tiene TIN
    for j, col in enumerate(OF):
        ho.formula(f'{col}{fila}', f'=ISNUMBER({col}{O["tin"]})', h='center', color='BBBBBB')
        nombre(wb, f'BANCO_OK_{j + 1}', 'Ofertas', f'{col}{fila}')
    ho.c(f'B{fila}', '(oferta rellenada)', font=E.f(8, False, 'BBBBBB'))
    ho.resaltar_si(f'D{O["coste"]}:G{O["coste"]}', f'AND(ISNUMBER(D{O["coste"]}),D{O["coste"]}=MIN($D{O["coste"]}:$G{O["coste"]}))', mezclar(E.acento, 0.45), bold=True)
    fila += 2
    ho.nota(f'B{fila}', AVISO + ' «Sin bonificar» usa un tipo constante (en las variables, el Euríbor del escenario más el diferencial sin bonificar).')
    ho.ws.merge_cells(f'B{fila}:G{fila + 1}')
    ho.ws.row_dimensions[fila].height = 30
    ho.ws.freeze_panes = 'D8'
    ho.proteger()

    # ------------------------------------------------------------ Cuadro mensual (motor)
    hm = Hoja(wb, 'Cuadro mensual', E, E.suave, {'A': 3, 'B': 7, 'C': 6})
    hm.cabecera('Cuadro mensual de las cuatro ofertas', 'Sistema francés con revisión anual del tipo. Lo usan las demás hojas: no hace falta tocarlo.', 'AA')
    hm.tabla_cab(R0 - 1, 2, ['Mes', 'Año'])
    for j in range(4):
        hm.tabla_cab(R0 - 2, base_col + j * 6, [f'=IF(OF_BANCO_{j + 1}="","Oferta {j + 1}",OF_BANCO_{j + 1})'] + [''] * 5, alto=18)
        hm.ws.merge_cells(f'{col_cm(j, 0)}{R0 - 2}:{col_cm(j, 5)}{R0 - 2}')
        hm.tabla_cab(R0 - 1, base_col + j * 6, ['TIN %', 'Cuota', 'Interés', 'Capital', 'Pendiente', 'Flujo (TAE)'])
        for k in range(6):
            hm.ws.column_dimensions[col_cm(j, k)].width = 11 if k else 7
    hm.formula(f'B{R0}', 0, h='center')
    for j in range(4):
        t, cu, it, ca, pe, fl = (col_cm(j, k) for k in range(6))
        hm.formula(f'{pe}{R0}', f'=IF(BANCO_OK_{j + 1},CAPITAL,0)', fmt='#,##0.00')
        hm.formula(f'{fl}{R0}', f'=IF(BANCO_OK_{j + 1},-(CAPITAL-GASTOS_{j + 1}),0)', fmt='#,##0.00')
    for m in range(1, MAXM + 1):
        r = R0 + m
        hm.formula(f'B{r}', m, h='center')
        hm.formula(f'C{r}', f'=ROUNDUP(B{r}/12,0)', h='center')
        for j in range(4):
            t, cu, it, ca, pe, fl = (col_cm(j, k) for k in range(6))
            tipo, tin, ai, dif = f'OF_TIPO_{j + 1}', f'OF_TIN_{j + 1}', f'OF_ANIOS_INI_{j + 1}', f'OF_DIF_{j + 1}'
            act = f'AND(BANCO_OK_{j + 1},B{r}<=MESES)'
            hm.formula(f'{t}{r}', f'=IF({act},IF(OR({tipo}="Fija",C{r}<={ai}),{tin},MAX(0,EURIBOR+{dif})),0)', fmt='0.000')
            prev_t = f'{t}{r - 1}' if m > 1 else '-1'
            prev_c = f'{cu}{r - 1}' if m > 1 else '0'
            hm.formula(f'{cu}{r}', f'=IF({act},IF(OR(B{r}=1,{t}{r}<>{prev_t}),PMT({t}{r}/1200,MESES-B{r}+1,-{pe}{r - 1}),{prev_c}),0)', fmt='#,##0.00')
            hm.formula(f'{it}{r}', f'={pe}{r - 1}*{t}{r}/1200', fmt='#,##0.00')
            hm.formula(f'{ca}{r}', f'=IF({act},IF(OR(B{r}=MESES,{cu}{r}-{it}{r}>{pe}{r - 1}),{pe}{r - 1},{cu}{r}-{it}{r}),0)', fmt='#,##0.00')
            hm.formula(f'{pe}{r}', f'={pe}{r - 1}-{ca}{r}', fmt='#,##0.00')
            hm.formula(f'{fl}{r}', f'=IF({act},{it}{r}+{ca}{r}+VINC_MES_{j + 1},0)', fmt='#,##0.00')
    nombre_rango(wb, 'CM_ANIO', 'Cuadro mensual', f'C{R0 + 1}:C{R0 + MAXM}')
    for j in range(4):
        for k, n in enumerate(['TIN', 'CUOTA', 'INT', 'CAP', 'PEND']):
            nombre_rango(wb, f'CMES_{n}_{j + 1}', 'Cuadro mensual', f'{col_cm(j, k)}{R0 + 1}:{col_cm(j, k)}{R0 + MAXM}')
    hm.ws.freeze_panes = f'D{R0}'
    hm.proteger()

    # ------------------------------------------------------------ Escenarios Euríbor
    he = Hoja(wb, 'Escenarios Euríbor', E, E.marca, {'A': 3, 'B': 36, 'C': 14, 'D': 18, 'E': 18, 'F': 18, 'G': 18, 'H': 3})
    fila = he.cabecera('Escenarios de Euríbor', 'Cuota después del periodo inicial con distintos Euríbor (constantes)', 'G')
    he.leyenda(fila); fila += 2
    he.tabla_cab(fila, 2, ['Euríbor', 'Cambio (puntos)', '=IF(OF_BANCO_1="","Oferta 1",OF_BANCO_1)', '=IF(OF_BANCO_2="","Oferta 2",OF_BANCO_2)', '=IF(OF_BANCO_3="","Oferta 3",OF_BANCO_3)', '=IF(OF_BANCO_4="","Oferta 4",OF_BANCO_4)'])
    fila += 1
    e0 = fila
    for k, d in enumerate([-1, 0, 1, 2, 3]):
        r = e0 + k
        he.entrada(f'C{r}', d, fmt='+0.0;-0.0;0')
        he.formula(f'B{r}', f'=EURIBOR_HOY+C{r}', fmt='0.000" %"', bold=True)
        for j, col in enumerate(OF):
            pend_ini = f'IF(OF_ANIOS_INI_{j + 1}=0,CAPITAL,INDEX(CMES_PEND_{j + 1},MIN(MESES,OF_ANIOS_INI_{j + 1}*12)))'
            he.formula(f'{col}{r}', f'=IF(NOT(BANCO_OK_{j + 1}),"",IF(OF_TIPO_{j + 1}="Fija",CUOTA1_{j + 1},PMT(MAX(0,B{r}+OF_DIF_{j + 1})/1200,MESES-OF_ANIOS_INI_{j + 1}*12,-{pend_ini})))', fmt=E.fmt_dinero)
    he.resaltar_si(f'B{e0}:G{e0 + 4}', f'$C{e0}=0', mezclar(E.acento, 0.6), bold=True)
    fila = e0 + 6
    he.etiqueta(f'B{fila}', 'Subida de la cuota si el Euríbor sube 2 puntos', bold=True)
    for j, col in enumerate(OF):
        he.formula(f'{col}{fila}', f'=IF({col}{e0 + 3}="","",{col}{e0 + 3}-{col}{e0 + 1})', fmt=E.fmt_dinero, bold=True)
    fila += 1
    he.etiqueta(f'B{fila}', 'Ingresos netos necesarios con +2 puntos')
    for j, col in enumerate(OF):
        he.formula(f'{col}{fila}', f'=IF({col}{e0 + 3}="","",({col}{e0 + 3}+VINC_MES_{j + 1})/(ESFUERZO/100))', fmt=E.fmt_dinero0)
    fila += 2
    he.nota(f'B{fila}', 'Las fijas no cambian. En las mixtas y variables, la cuota se recalcula cada año con el Euríbor del momento más el diferencial; aquí se supone el mismo Euríbor durante todo el resto del préstamo. Nadie puede predecir el Euríbor: úsalo para ver si podrías pagar en un escenario malo.')
    he.ws.merge_cells(f'B{fila}:G{fila + 1}')
    he.ws.row_dimensions[fila].height = 30
    he.proteger()

    # ------------------------------------------------------------ Amortizar antes
    ha = Hoja(wb, 'Amortizar antes', E, E.marca, {'A': 3, 'B': 52, 'C': 18, 'D': 18, 'E': 52})
    fila = ha.cabecera('Amortizar antes', 'Si un día te sobra dinero: ¿reducir cuota o reducir plazo?', 'E')
    ha.leyenda(fila); fila += 2
    A = {}
    def la(clave, texto, f_, fmt=E.fmt_dinero, entrada=False, bold=False, res=False, nota_=''):
        nonlocal fila
        ha.etiqueta(f'B{fila}', texto, bold=bold or entrada)
        if entrada:
            ha.entrada(f'C{fila}', f_, fmt=fmt)
        elif res:
            ha.resultado(f'C{fila}', f_, fmt=fmt)
        else:
            ha.formula(f'C{fila}', f_, fmt=fmt, bold=bold)
        if nota_:
            ha.nota(f'D{fila}', nota_)
            ha.ws.merge_cells(f'D{fila}:E{fila}')
        A[clave] = f'C{fila}'
        nombre(wb, 'AA_' + clave.upper(), 'Amortizar antes', f'C{fila}')
        fila += 1
    la('oferta', 'Oferta (1 a 4)', 1, fmt='0', entrada=True)
    ha.lista(A['oferta'], ['1', '2', '3', '4'])
    la('anio', 'Amortizas al final del año', 5, fmt='0', entrada=True)
    la('importe', 'Importe que amortizas (€)', 20000, entrada=True)
    la('opcion', 'Si es variable o mixta, opción de comisión del contrato', 'A', fmt=None, entrada=True, nota_='A: 0,25 % los 3 primeros años; B: 0,15 % los 5 primeros. Mira tu escritura o la FEIN.')
    ha.lista(A['opcion'], ['A', 'B'])
    fila += 1
    ha.seccion(fila, 'Situación en ese momento', 'B', 'E'); fila += 1
    sel = lambda pre: f'CHOOSE({A["oferta"]},{",".join(f"{pre.format(j + 1)}" for j in range(4))})'
    mes = f'MIN(MESES,{A["anio"]}*12)'
    la('pendiente', 'Capital pendiente', f'=INDEX(CHOOSE({A["oferta"]},CMES_PEND_1,CMES_PEND_2,CMES_PEND_3,CMES_PEND_4),{mes})')
    la('tin', 'TIN del año siguiente (%)', f'=INDEX(CHOOSE({A["oferta"]},CMES_TIN_1,CMES_TIN_2,CMES_TIN_3,CMES_TIN_4),MIN(MESES,{mes}+1))', fmt='0.000')
    la('n', 'Meses que quedan', f'=MESES-{mes}', fmt='0')
    la('tipo', 'Tipo de la oferta', '=' + sel('OF_TIPO_{}'), fmt=None)
    i = f'{A["tin"]}/1200'
    la('cuotaantes', 'Cuota si no amortizas', f'=PMT({i},{A["n"]},-{A["pendiente"]})')
    la('intantes', 'Intereses que te quedan si no amortizas', f'={A["cuotaantes"]}*{A["n"]}-{A["pendiente"]}')
    la('nuevo', 'Capital pendiente después de amortizar', f'=MAX(0,{A["pendiente"]}-{A["importe"]})')
    fila += 1
    ha.seccion(fila, 'Reducir cuota (mismo plazo)', 'B', 'E'); fila += 1
    la('cuotacuota', 'Cuota nueva', f'=PMT({i},{A["n"]},-{A["nuevo"]})', bold=True)
    la('intcuota', 'Intereses que te quedan', f'={A["cuotacuota"]}*{A["n"]}-{A["nuevo"]}')
    la('ahorrocuota', 'Ahorro de intereses', f'={A["intantes"]}-{A["intcuota"]}', res=True)
    fila += 1
    ha.seccion(fila, 'Reducir plazo (misma cuota)', 'B', 'E'); fila += 1
    la('mesesplazo', 'Meses que quedan', f'=IF({A["nuevo"]}<=0,0,ROUNDUP(NPER({i},-{A["cuotaantes"]},{A["nuevo"]})-1E-9,0))', fmt='0', bold=True)
    la('mesesmenos', 'Meses que te ahorras', f'={A["n"]}-{A["mesesplazo"]}', fmt='0')
    la('intplazo', 'Intereses que te quedan', f'=IF({A["mesesplazo"]}=0,0,PMT({i},{A["mesesplazo"]},-{A["nuevo"]})*{A["mesesplazo"]}-{A["nuevo"]})')
    la('ahorroplazo', 'Ahorro de intereses', f'={A["intantes"]}-{A["intplazo"]}', res=True)
    fila += 1
    ha.seccion(fila, 'Comisión', 'B', 'E'); fila += 1
    tope = (f'IF({A["tipo"]}="Fija",IF({A["anio"]}<10,COM_FIJO_10,COM_FIJO_RESTO),IF({A["opcion"]}="B",IF({A["anio"]}<COM_VAR_B_ANIOS,COM_VAR_B,COM_VAR_DESPUES),IF({A["anio"]}<COM_VAR_A_ANIOS,COM_VAR_A,COM_VAR_DESPUES)))')
    la('topepct', 'Tope legal de la comisión (%)', '=' + tope, fmt='0.00')
    la('comision', 'Comisión máxima (€)', f'=MIN({A["importe"]},{A["pendiente"]})*{A["topepct"]}/100', nota_='La ley la limita también a la pérdida financiera real del banco (Ley 5/2019, art. 23): puede ser menor.')
    la('neto', 'Ahorro neto reduciendo plazo (tras la comisión)', f'={A["ahorroplazo"]}-{A["comision"]}', res=True)
    fila += 1
    ha.nota(f'B{fila}', 'Reducir plazo casi siempre ahorra más intereses; reducir cuota te da más margen cada mes. Calculado con el TIN del año siguiente constante hasta el final. ' + AVISO)
    ha.ws.merge_cells(f'B{fila}:E{fila + 1}')
    ha.ws.row_dimensions[fila].height = 30
    ha.proteger()

    # ------------------------------------------------------------ Cuadro anual
    hy = Hoja(wb, 'Cuadro anual', E, E.marca, {'A': 3, 'B': 9, 'C': 12, 'D': 15, 'E': 16, 'F': 16, 'G': 16, 'H': 17})
    fila = hy.cabecera('Cuadro de amortización por años', '="Oferta: "&CHOOSE(CA_OFERTA,OF_BANCO_1,OF_BANCO_2,OF_BANCO_3,OF_BANCO_4)', 'H')
    hy.etiqueta(f'B{fila}', 'Oferta', bold=True)
    hy.entrada(f'C{fila}', 1, fmt='0', h='center')
    hy.lista(f'C{fila}', ['1', '2', '3', '4'])
    nombre(wb, 'CA_OFERTA', 'Cuadro anual', f'C{fila}')
    hy.nota(f'D{fila}', 'Elige 1, 2, 3 o 4')
    fila += 2
    hy.tabla_cab(fila, 2, ['Año', 'TIN %', 'Cuota media', 'Pagado en el año', 'Intereses', 'Capital', 'Pendiente al final'])
    fila += 1
    y0 = fila
    selc = lambda n: f'CHOOSE(CA_OFERTA,CMES_{n}_1,CMES_{n}_2,CMES_{n}_3,CMES_{n}_4)'
    for k in range(1, 41):
        r = y0 + k - 1
        hy.formula(f'B{r}', f'=IF({k}*12-11>MESES,"",{k})', h='center')
        hy.formula(f'C{r}', f'=IF(B{r}="","",INDEX({selc("TIN")},{k}*12-11))', fmt='0.000')
        hy.formula(f'E{r}', f'=IF(B{r}="","",SUMIFS({selc("INT")},CM_ANIO,{k})+SUMIFS({selc("CAP")},CM_ANIO,{k}))', fmt=E.fmt_dinero)
        hy.formula(f'D{r}', f'=IF(B{r}="","",E{r}/MIN(12,MESES-({k}-1)*12))', fmt=E.fmt_dinero)
        hy.formula(f'F{r}', f'=IF(B{r}="","",SUMIFS({selc("INT")},CM_ANIO,{k}))', fmt=E.fmt_dinero)
        hy.formula(f'G{r}', f'=IF(B{r}="","",SUMIFS({selc("CAP")},CM_ANIO,{k}))', fmt=E.fmt_dinero)
        hy.formula(f'H{r}', f'=IF(B{r}="","",INDEX({selc("PEND")},MIN(MESES,{k}*12)))', fmt=E.fmt_dinero)
        if k in (1, 10):
            nombre(wb, f'CA_INT_{k}', 'Cuadro anual', f'F{r}')
            nombre(wb, f'CA_PEND_{k}', 'Cuadro anual', f'H{r}')
    y1 = y0 + 39
    hy.formula(f'B{y1 + 1}', 'Total', bold=True, h='center')
    for col in 'EFG':
        hy.formula(f'{col}{y1 + 1}', f'=SUM({col}{y0}:{col}{y1})', fmt=E.fmt_dinero, bold=True)
    hy.ws.freeze_panes = f'B{y0}'
    hy.proteger()

    # ------------------------------------------------------------ Comparativa (portada de resultados)
    hc = Hoja(wb, 'Comparativa', E, E.acento, {'A': 3, 'B': 44, 'C': 18, 'D': 18, 'E': 18, 'F': 18, 'G': 3})
    fila = hc.cabecera('Comparativa', '="Préstamo de "&FIXED(CAPITAL,0)&" € a "&PLAZO&" años · Euríbor "&EURIBOR&" %"', 'F')
    fila += 1
    hc.tabla_cab(fila, 2, ['', '=IF(OF_BANCO_1="","Oferta 1",OF_BANCO_1)', '=IF(OF_BANCO_2="","Oferta 2",OF_BANCO_2)', '=IF(OF_BANCO_3="","Oferta 3",OF_BANCO_3)', '=IF(OF_BANCO_4="","Oferta 4",OF_BANCO_4)'])
    fila += 1
    k0 = fila
    for clave, texto, fmt, fuerte in [('tipo', 'Tipo', None, False), ('cuota1', 'Cuota inicial', E.fmt_dinero, True), ('cuota2', 'Cuota después del periodo inicial', E.fmt_dinero, False),
                                      ('vincmes', 'Seguros y vinculados al mes', E.fmt_dinero, False), ('tae', 'TAE aproximada (con vinculaciones)', '0.00" %"', True),
                                      ('intereses', 'Intereses totales', E.fmt_dinero0, False), ('vinctot', 'Vinculaciones totales', E.fmt_dinero0, False),
                                      ('gastos', 'Gastos iniciales', E.fmt_dinero0, False), ('coste', 'COSTE TOTAL', E.fmt_dinero0, True),
                                      ('compensa', '¿Compensa bonificarte?', None, False), ('rank', 'Posición', '0"º"', True)]:
        hc.etiqueta(f'B{fila}', texto, bold=fuerte)
        for j, col in enumerate('CDEF'):
            f_ = f'=OF_TIPO_{j + 1}' if clave == 'tipo' else f'={clave.upper()}_{j + 1}'
            if clave == 'tipo':
                f_ = f'=IF(BANCO_OK_{j + 1},OF_TIPO_{j + 1}&IF(OF_TIPO_{j + 1}="Fija"," "&OF_TIN_{j + 1}&" %"," · "&OF_ANIOS_INI_{j + 1}&IF(OF_ANIOS_INI_{j + 1}=1," año"," años")&" + Euríbor + "&OF_DIF_{j + 1}),"")'
            (hc.resultado if fuerte else hc.formula)(f'{col}{fila}', f_, fmt=fmt, **({} if fuerte else {}))
            if clave in ('tipo', 'compensa'):
                hc.ws[f'{col}{fila}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if clave in ('compensa', 'tipo'):
            hc.ws.row_dimensions[fila].height = 30
        if clave == 'coste':
            fcoste = fila
        fila += 1
    hc.resaltar_si(f'C{fcoste}:F{fcoste}', f'AND(ISNUMBER(C{fcoste}),C{fcoste}=MIN($C{fcoste}:$F{fcoste}))', mezclar(E.acento, 0.45), bold=True)
    fila += 1
    hc.formula(f'B{fila}', '=IF(COUNT(C' + str(fcoste) + ':F' + str(fcoste) + ')<2,"Rellena al menos dos ofertas para compararlas.","La más barata en todo el plazo es "&INDEX(C' + str(k0 - 1) + ':F' + str(k0 - 1) + ',MATCH(MIN(C' + str(fcoste) + ':F' + str(fcoste) + '),C' + str(fcoste) + ':F' + str(fcoste) + ',0))&": "&FIXED(MIN(C' + str(fcoste) + ':F' + str(fcoste) + '),0)&" € de coste total.")', h='left', bold=True, color=E.pri)
    hc.ws.merge_cells(f'B{fila}:F{fila}')
    fila += 2
    graf = BarChart()
    graf.type = 'col'
    graf.title = 'Coste total de cada oferta (€)'
    graf.style = 10
    graf.y_axis.title = None
    graf.y_axis.numFmt = '#,##0'
    graf.y_axis.majorGridlines = None
    graf.y_axis.scaling.min = 0
    graf.y_axis.delete = False
    graf.x_axis.delete = False
    graf.legend = None
    datos = Reference(hc.ws, min_col=3, max_col=6, min_row=fcoste, max_row=fcoste)
    cats = Reference(hc.ws, min_col=3, max_col=6, min_row=k0 - 1, max_row=k0 - 1)
    graf.add_data(datos, from_rows=True, titles_from_data=False)
    graf.set_categories(cats)
    graf.series[0].graphicalProperties.solidFill = E.pri
    graf.series[0].graphicalProperties.line.solidFill = E.pri
    graf.dataLabels = DataLabelList()
    graf.dataLabels.showVal = True
    for k in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent', 'showBubbleSize'):
        setattr(graf.dataLabels, k, False)
    graf.varyColors = False
    graf.height, graf.width = 7.5, 17
    hc.ws.add_chart(graf, f'B{fila}')
    hc.nota(f'B{fila + 16}', AVISO, italic=True)
    hc.ws.merge_cells(f'B{fila + 16}:F{fila + 16}')
    hc.proteger()

    orden = ['Portada', 'Instrucciones', 'Comparativa', 'Tu hipoteca', 'Ofertas', 'Escenarios Euríbor', 'Amortizar antes', 'Cuadro anual', 'Cuadro mensual', 'Parámetros']
    wb._sheets = [wb[n] for n in orden]
    return wb


# ---------------------------------------------------------------- PDF: preguntas para el banco

def guia_html() -> str:
    E = Estilo('vivienda')
    url = url_publica('vivienda')
    eu = ultimo_euribor()
    C = V['comision_amortizacion_anticipada']
    cuerpo = portada_pdf(E, 'Preguntas para el banco antes de firmar la hipoteca', 'Lo que hay que pedir, preguntar y comparar para negociar el tipo, las vinculaciones y las comisiones con datos.',
                         ['FEIN y FiAE', 'TAE con vinculaciones', 'Seguros', 'Comisiones', 'Euríbor', 'Notaría'],
                         f'{E.nombre} · {url.replace("https://", "")} · hola@numimates.com<br>Información orientativa: no es asesoramiento financiero. Contenido preparado con ayuda de IA y revisado con fuentes oficiales (Ley 5/2019, Banco de España).')
    q = lambda t, d='': f'<li><strong>{t}</strong>{" " + d if d else ""}</li>'
    cuerpo += f"""
<h2>1. Lo que te tienen que dar por escrito</h2>
<ul class="check">{q('La FEIN (Ficha Europea de Información Normalizada) con tus condiciones personalizadas.', 'Es gratuita y te obliga a pedirla a cada banco: sin FEIN no hay comparación.')}
{q('La FiAE (Ficha de Advertencias Estandarizadas)', 'con las cláusulas importantes: tipo mínimo, vencimiento anticipado, comisiones.')}
{q('El proyecto de contrato y el coste de los seguros', 'que te exigen o que bonifican el tipo.')}
{q('Al menos 10 días naturales antes de firmar', '(Ley 5/2019, art. 14). Después irás al notario a un acta previa gratuita para que te explique el contrato (art. 15).')}</ul>

<h2>2. El tipo de interés</h2>
<ol class="pasos">{q('¿Cuál es el TIN sin ninguna bonificación y cuál con todas?', 'Apunta los dos: el comparador te dice si bonificarte compensa.')}
{q('Si es mixta: ¿cuántos años a tipo fijo y qué diferencial después?', 'El diferencial es para siempre; una décima menos vale más que un mes gratis.')}
{q('Si es variable: ¿con qué Euríbor se revisa y cada cuánto?', f'Normalmente anual con el dato de dos meses antes. El último Euríbor publicado es {str(eu["valor"]).replace(".", ",")} % ({eu["mes"]}).')}
{q('¿La TAE de la FEIN incluye los seguros y productos vinculados?', 'Pide la TAE con todas las bonificaciones y compara la del comparador.')}
{q('¿Qué pasa si dejo de cumplir una bonificación?', '¿Sube el tipo en esa revisión, desde cuándo y cuánto? ¿Hay periodo de gracia?')}</ol>

<h2>3. Vinculaciones y seguros</h2>
<ol class="pasos">{q('¿Qué productos son obligatorios y cuáles solo bonifican?', 'La venta vinculada está prohibida salvo excepciones (Ley 5/2019, art. 17): el banco puede exigir un seguro de daños del inmueble, pero puedes contratarlo con otra compañía.')}
{q('¿Cuánto cuesta el seguro de vida el primer año y cómo sube con la edad?', 'Pide la tabla de primas por edades o el compromiso de prima fija.')}
{q('¿Prima anual o prima única?', 'La prima única financiada encarece el préstamo: evítala.')}
{q('¿Puedo cambiar de seguro al año siguiente sin perder la bonificación?', 'Pregunta qué condiciones debe cumplir el seguro alternativo.')}
{q('¿Cuánto me cuesta la domiciliación de la nómina, la tarjeta, el plan de pensiones?', 'Ponlo como «otros productos vinculados» al mes.')}</ol>

<h2>4. Comisiones y gastos</h2>
<table><tr><th>Concepto</th><th>Qué dice la ley</th><th>Pregunta</th></tr>
<tr><td>Apertura</td><td>Debe englobar todos los gastos de estudio y tramitación</td><td>¿Hay comisión de apertura? ¿Se puede quitar?</td></tr>
<tr><td>Amortización anticipada (fija)</td><td>Máximo {str(C['fijo']['primeros_10_anios_pct']).replace('.', ',')} % los 10 primeros años y {str(C['fijo']['resto_pct']).replace('.', ',')} % después</td><td>¿Qué porcentaje y desde qué año?</td></tr>
<tr><td>Amortización anticipada (variable)</td><td>{str(C['variable']['opcion_a']['tope_pct']).replace('.', ',')} % los {C['variable']['opcion_a']['durante_anios']} primeros años o {str(C['variable']['opcion_b']['tope_pct']).replace('.', ',')} % los {C['variable']['opcion_b']['durante_anios']} primeros; después, {C['variable']['despues']} %</td><td>¿Qué opción tiene el contrato?</td></tr>
<tr><td>Notaría, registro, gestoría, AJD de la hipoteca</td><td>Los paga el banco</td><td>¿Me cobran algo de esto? (no deberían)</td></tr>
<tr><td>Tasación</td><td>La pagas tú, pero puedes aportar una tasación válida de otra tasadora</td><td>¿Aceptan mi tasación?</td></tr></table>
<p class="fuentes">Ley 5/2019, reguladora de los contratos de crédito inmobiliario: arts. 14, 15, 17, 23 y 24. Topes de comisión: {fuente_de(C)[0]}.</p>

<h2>5. Para negociar</h2>
<ul><li><strong>Lleva la mejor FEIN de otro banco</strong> impresa: es el mejor argumento. Pregunta si la igualan o la mejoran en diferencial o en vinculaciones.</li>
<li><strong>Negocia el diferencial y el tipo sin bonificar</strong>, no solo el bonificado: si un día dejas la nómina, eso es lo que pagarás.</li>
<li><strong>Pide quitar productos</strong> que no necesitas a cambio de una décima más: con el comparador verás si te sale a cuenta.</li>
<li><strong>Escenario malo:</strong> antes de firmar una variable o mixta, mira en «Escenarios Euríbor» la cuota con dos puntos más. Si no podrías pagarla, busca más plazo fijo.</li></ul>
<p class="legal">Información orientativa: no es asesoramiento financiero. Contenido preparado con ayuda de inteligencia artificial y revisado con las fuentes oficiales citadas. Parámetros de params/es/{ANIO}/vivienda.json a {V['actualizado']}. Dudas sobre el producto: hola@numimates.com.</p>"""
    return html_pdf(E, 'Preguntas para el banco antes de firmar la hipoteca', cuerpo)


def main():
    out = SALIDA / ID
    guardar(construir(), out / 'comparador-hipotecas.xlsx')
    escribir(out / 'fuente' / 'preguntas-banco-hipoteca.html', guia_html())
    if '--demo' in sys.argv:
        guardar(construir(demo=True), out / 'demo' / 'comparador-hipotecas-demo.xlsx')
    print('Comparador de hipotecas: listo en', out)


if __name__ == '__main__':
    main()
