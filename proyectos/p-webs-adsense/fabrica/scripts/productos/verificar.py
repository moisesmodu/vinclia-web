"""Rellena un Excel de producto con unos datos, lo recalcula con LibreOffice y devuelve los resultados (JSON).

    python3 scripts/productos/verificar.py casos.json   → imprime [{"id": ..., "valores": {nombre: valor}}]

casos.json: [{"id": "caso1", "xlsx": "productos/x/x.xlsx", "poner": {"NOMBRE_DEFINIDO" | "Hoja!B8": valor}, "leer": ["NOMBRE", "Hoja!C9"]}]
Las fechas van como "AAAA-MM-DD". Lo usa tests/productos.test.ts para comparar las fórmulas con los motores de la web.
Necesita LibreOffice (soffice). Todos los casos de una llamada se recalculan con una sola ejecución de soffice.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

FABRICA = Path(__file__).resolve().parents[2]
ISO = re.compile(r'^\d{4}-\d{2}-\d{2}$')


def ref(wb, clave: str):
    if '!' in clave:
        hoja, celda = clave.rsplit('!', 1)
        return wb[hoja.strip("'")], celda.replace('$', '')
    dn = wb.defined_names[clave]
    hoja, celda = dn.attr_text.rsplit('!', 1)
    return wb[hoja.strip("'")], celda.replace('$', '')


def valor(v):
    if isinstance(v, str) and ISO.match(v):
        return date.fromisoformat(v)
    return v


def salida(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return v


def main():
    casos = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    tmp = Path(tempfile.mkdtemp(prefix='verificar-'))
    try:
        entradas = []
        for k, c in enumerate(casos):
            wb = openpyxl.load_workbook(FABRICA / c['xlsx'])
            for clave, v in c.get('poner', {}).items():
                ws, celda = ref(wb, clave)
                ws[celda] = valor(v)
            p = tmp / f'caso{k}.xlsx'
            wb.save(p)
            entradas.append(p)
        perfil = tmp / 'perfil'
        r = subprocess.run(['soffice', f'-env:UserInstallation=file://{perfil}', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(tmp / 'out'), *map(str, entradas)],
                           capture_output=True, text=True, timeout=300)
        res = []
        for k, c in enumerate(casos):
            out = tmp / 'out' / f'caso{k}.xlsx'
            if not out.exists():
                raise SystemExit(f'LibreOffice no recalculó {c["id"]}: {r.stderr}')
            wb = openpyxl.load_workbook(out, data_only=True)
            vals = {}
            for clave in c['leer']:
                ws, celda = ref(wb, clave)
                vals[clave] = salida(ws[celda].value)
            res.append({'id': c['id'], 'valores': vals})
        print(json.dumps(res, ensure_ascii=False))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    main()
